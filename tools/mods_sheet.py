"""Makes Ahrivaea_Mods.xlsx (every mod, what it does, and which worlds use it) from mods.json.
Usage: python3 tools/mods_sheet.py [output.xlsx]"""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

mods = json.load(open('mods.json'))
version = re.search(r'^version=(.+)$', open('gradle.properties').read(), re.M).group(1).strip()
out = sys.argv[1] if len(sys.argv) > 1 else 'Ahrivaea_Mods.xlsx'
F = 'Arial'; f = lambda **k: Font(name=F, **k)
thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr = PatternFill('solid', fgColor='262626'); green = PatternFill('solid', fgColor='C6EFCE')
wrap = Alignment(wrap_text=True, vertical='top')

def header(ws, row, cols, widths):
    for j, (h, w) in enumerate(zip(cols, widths)):
        c = ws.cell(row, 1 + j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hdr; c.border = B
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[c.column_letter].width = w

wb = Workbook(); ws = wb.active; ws.title = 'Mods'
ws['A1'] = 'Ahrivaea Mods'; ws['A1'].font = f(bold=True, size=14)
ws['A2'] = f'Version {version}. Each mod is its own file: put the files a world needs in its server\'s mods folder (and in each player\'s). Download them from the "latest" release on GitHub.'
ws['A2'].font = f(italic=True, size=9, color='595959')
header(ws, 4, ['Mod', 'File', 'What it does', 'Features', 'Commands', 'Players need it installed?', 'Status', 'Added', 'More detail'],
       [24, 30, 40, 70, 34, 22, 22, 11, 22])
for i, m in enumerate(mods['mods']):
    r = 5 + i
    row = [m['name'], f"ahrivaea-{m['key']}-{version}.jar", m['summary'], '\n'.join('• ' + x for x in m['features']), m['commands'],
           m['players_install'], m['status'], m['added'], m['docs']]
    for j, v in enumerate(row):
        c = ws.cell(r, 1 + j, v); c.font = f(bold=(j == 0)); c.border = B; c.alignment = wrap
ws.freeze_panes = 'B5'

w2 = wb.create_sheet('Worlds')
w2['A1'] = 'Which mods each world uses'; w2['A1'].font = f(bold=True, size=14)
w2['A2'] = 'Put ✓ where a world uses a mod. Add a world by filling in a new column. Those files go in that world\'s server mods folder (and each player\'s).'
w2['A2'].font = f(italic=True, size=9, color='595959')
worlds = mods['worlds'] + [{'name': '(new world)', 'notes': '', 'mods': []}] * 2
header(w2, 4, ['Mod', 'File'] + [w['name'] for w in worlds], [26, 30] + [22] * len(worlds))
dv = DataValidation(type='list', formula1='"✓"', allow_blank=True); w2.add_data_validation(dv)
for i, m in enumerate(mods['mods']):
    r = 5 + i
    w2.cell(r, 1, m['name']).font = f(bold=True); w2.cell(r, 2, f"ahrivaea-{m['key']}-{version}.jar").font = f()
    for j, w in enumerate(worlds):
        c = w2.cell(r, 3 + j, '✓' if m['key'] in w['mods'] else None)
        c.font = f(color='006100', bold=True); c.alignment = Alignment(horizontal='center'); dv.add(c)
    for j in range(len(worlds) + 2): w2.cell(r, 1 + j).border = B
r = 5 + len(mods['mods'])
w2.cell(r, 1, 'Notes').font = f(italic=True)
for j, w in enumerate(worlds):
    c = w2.cell(r, 3 + j, w['notes']); c.font = f(italic=True, size=9); c.alignment = wrap
from openpyxl.formatting.rule import CellIsRule
w2.conditional_formatting.add(f"C5:{w2.cell(4 + len(mods['mods']), 2 + len(worlds)).coordinate}", CellIsRule(operator='equal', formula=['"✓"'], fill=green))

w3 = wb.create_sheet('Adding & removing')
lines = [('Adding or removing a mod', True), ('', False),
 ('To add a mod to a world: put its .jar file in the server\'s mods folder, and in the mods folder of every player who joins it. Restart the server.', False),
 ('To remove one: take its .jar out of the same places and restart.', False),
 ('Every mod also needs Fabric Loader 0.19.5+ (installed with the Fabric installer) and Fabric API: put the fabric-api-...jar from the same', False),
 ('download page in the mods folder too. Without it Minecraft says the mods "require any version of fabric-api, which is missing".', False), ('', False),
 ('Each mod works on its own: any mix of them can be used in a world.', False), ('', False),
 ('Before removing a mod from a world people have played in:', True),
 ('  • Its items disappear from inventories and chests (e.g. removing Stellar Gear deletes Stellar tools and armor).', False),
 ('  • Its creatures disappear (Creatures), or its towers stay as plain blocks but stop working (Tower Challenges).', False),
 ('  • Adding it back later does not bring removed items back. Make a backup of the world first.', False), ('', False),
 ('New mods: whenever a new mod is made, it is added to this sheet and gets its own file.', False)]
for i, (t, b) in enumerate(lines):
    c = w3.cell(1 + i, 1, t); c.font = f(bold=b, size=14 if (b and i == 0) else 10)
w3.column_dimensions['A'].width = 130
wb.save(out)
print('wrote', out)
