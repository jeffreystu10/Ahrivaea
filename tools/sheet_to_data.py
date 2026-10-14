"""Turns Tower_Challenges.xlsx into the mod's challenge files and challenge_order.json.
Usage: python3 tools/sheet_to_data.py Tower_Challenges.xlsx [path/to/minecraft_extracted]  (the second
argument, a folder with assets/minecraft/items and data/minecraft/tags/item, checks every item id)."""
import json, os, re, sys
from openpyxl import load_workbook

D = 'src/main/resources/data/ahrivaea/'
sheet = load_workbook(sys.argv[1], data_only=True)['Towers']
mc = sys.argv[2] if len(sys.argv) > 2 else None
problems = []

def item_ok(i):
    if not mc: return True
    if i.startswith('#'): return os.path.exists(f"{mc}/data/minecraft/tags/item/{i[1:]}.json")
    return os.path.exists(f"{mc}/assets/minecraft/items/{i}.json")

def full(i):
    return ('#minecraft:' + i[1:]) if i.startswith('#') and ':' not in i else (('minecraft:' + i) if ':' not in i else i)

def parse_items(text, row):
    out = {}
    for part in filter(None, (p.strip() for p in str(text or '').split(','))):
        m = re.fullmatch(r'(#?[a-z0-9_:/]+)\s*=\s*(\d+)', part)
        if not m: problems.append(f"row {row}: can't read '{part}' (use item = number)"); continue
        if not item_ok(m.group(1)): problems.append(f"row {row}: unknown item {m.group(1)}")
        out[full(m.group(1))] = int(m.group(2))
    return out

def parse_reward(text, row):
    out = []
    for part in filter(None, (p.strip() for p in str(text or '').split(','))):
        m = re.fullmatch(r'(\d+)\s*x?\s+([a-z0-9_:]+)', part)
        if not m: problems.append(f"row {row}: can't read reward '{part}' (use number item)"); continue
        if not item_ok(m.group(2)): problems.append(f"row {row}: unknown reward item {m.group(2)}")
        out.append({"item": full(m.group(2)), "count": int(m.group(1))})
    return out

head = [c.value for c in sheet[4]]
col = {h: i for i, h in enumerate(head) if h}
order, seen = [], set()
for r in range(5, sheet.max_row + 1):
    v = [c.value for c in sheet[r]]
    get = lambda h: v[col[h]] if h in col and col[h] < len(v) else None
    if not get('Challenge name') or not get('What to bring (item = how much it counts)'): continue
    cid = str(get('ID') or re.sub(r'[^a-z0-9]+', '_', str(get('Challenge name')).lower()).strip('_'))
    if cid in seen: problems.append(f"row {r}: ID {cid} is used twice"); continue
    seen.add(cid)
    rewards = []
    for rank, h in ((1, '1st place reward'), (2, '2nd place reward'), (3, '3rd place reward'), (0, 'Everyone who helped')):
        items = parse_reward(get(h), r)
        if items: rewards.append({"rank": rank, "items": items})
    d = {"name": get('Challenge name'), "unit": get('Unit word') or 'items', "items": parse_items(get('What to bring (item = how much it counts)'), r),
         "goal": int(get('Goal')), "stages": int(get('Stages') or 4),
         "ruined": "ahrivaea:tower/ruined", "finished": "ahrivaea:tower/finished", "chest": [5, 1, 5], "rewards": rewards}
    json.dump(d, open(D + f'challenge/{cid}.json', 'w'), indent=2)
    order.append(cid)
for f in os.listdir(D + 'challenge'):
    if f.endswith('.json') and f[:-5] not in seen: os.remove(D + 'challenge/' + f)
json.dump({"towers": order}, open(D + 'challenge_order.json', 'w'), indent=2)
print(len(order), 'challenges written')
for p in problems: print('PROBLEM', p)
sys.exit(1 if problems else 0)
